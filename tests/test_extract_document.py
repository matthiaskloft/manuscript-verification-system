"""`extract_document`: one read of a document, for everything a check needs.

Phase 4 of docs/plans/plan-in-text-citation-parsing.md re-plumbs extraction so that the
references, the body text and GROBID's in-text citations all come out of a single pass.
Before this, a caller wanting both halves made one GROBID round-trip for the references
and then parsed the PDF again for the artifact.

The claim that matters most here is the one the plan's first Success Criterion makes:
reference checking produces exactly what it produced before, whatever citation matching
does or fails to do. That is checked by running the new entry point and the old one over
the same fixtures and comparing entry for entry — not by trusting that a shared
implementation stayed shared.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from refcheck.extraction import document
from refcheck.extraction.document import (
    NoBibliographySectionError,
    build_document_artifact,
    extract_document,
    extract_references,
)
from refcheck.extraction.engine_status import ENGINE_ANCHOR, ENGINE_GROBID
from refcheck.extraction.grobid import CitationContext, GrobidDocument, GrobidUnavailableError
from refcheck.extraction.tier0 import RawReferenceEntry

FIXTURES = Path(__file__).parent / "fixtures" / "synthetic"
MANIFEST = json.loads((FIXTURES / "manifest.json").read_text(encoding="utf-8"))
PDFS = [FIXTURES / doc["pdf"] for doc in MANIFEST["documents"]] + [
    FIXTURES / "corruption_variants" / name
    for name in ("corrupt_running_header.pdf", "corrupt_two_column.pdf", "reference_page_break.pdf")
]


@pytest.fixture(autouse=True)
def _no_grobid(monkeypatch):
    """Tier 0 everywhere unless a test says otherwise — a suite whose result depends on
    whether a container happens to be running locally measures nothing."""
    monkeypatch.setattr(document, "is_grobid_available", lambda: False)


@pytest.mark.parametrize("path", PDFS, ids=lambda p: p.name)
def test_the_references_are_the_ones_extract_references_produces(path):
    """The Success Criterion, checked directly rather than argued from shared code."""
    assert extract_document(path).references == extract_references(path)


@pytest.mark.parametrize("path", PDFS[:2], ids=lambda p: p.name)
def test_the_body_is_the_one_build_document_artifact_produces(path):
    """The other half of the same promise, for the artifact: reading it alongside the
    references must not change what it contains."""
    artifact = extract_document(path).artifact

    assert artifact.body_text == build_document_artifact(path).body_text
    assert artifact.parser == "pymupdf4llm"


def test_a_docx_is_read_the_same_way(tmp_path):
    import docx

    written = docx.Document()
    for paragraph in [
        "An opening paragraph citing Doe (2020) in passing.",
        "References",
        "Doe, J. (2020). A title.",
        "Roe, R. (2019). Another title.",
    ]:
        written.add_paragraph(paragraph)
    path = tmp_path / "manuscript.docx"
    written.save(path)

    extracted = extract_document(path)

    assert extracted.references == extract_references(path)
    assert extracted.artifact.parser == "python-docx"
    assert "Doe (2020)" in extracted.artifact.body_text
    assert extracted.citations == ()


def test_the_pdf_is_parsed_once(monkeypatch):
    """The cost this exists to remove. Phase 1 recorded the double parse as something
    Phase 4 would otherwise pay on every check, once detection became unconditional."""
    calls = []
    original = document._extract_pdf_markdown
    monkeypatch.setattr(
        document, "_extract_pdf_markdown", lambda path: (calls.append(path), original(path))[1]
    )

    extract_document(PDFS[0])

    assert len(calls) == 1


def test_the_bibliographys_style_profile_comes_back_with_it():
    """`match_citations` needs it to narrow which marker families to look for, and it is
    computed from the same bibliography slice the references were split out of — a second
    notion of where the reference list is would be the divergence signals.py is about."""
    extracted = extract_document(FIXTURES / MANIFEST["documents"][0]["pdf"])

    assert extracted.profile is not None
    assert extracted.profile.paren_year_matches > 0


# --------------------------------------------------------------------------------------
# The GROBID path
# --------------------------------------------------------------------------------------


def _pdf(tmp_path: Path) -> Path:
    path = tmp_path / "manuscript.pdf"
    path.write_bytes(b"%PDF-1.4 fake")
    return path


_SOURCE = "Body text citing [1] here.\n\nReferences\n\n[1] Doe, J. (2020). A title.\n\n[2] Roe, R. (2019). Another."


def test_grobid_references_and_citations_come_from_one_call(monkeypatch, tmp_path):
    calls = []

    def fake_grobid(path):
        calls.append(path)
        return GrobidDocument(
            references=[RawReferenceEntry(1, "Doe, J. (2020) A Title", title="A Title", source_id="b0")],
            citations=[CitationContext(1, "b0", "[1]", "a paragraph", 0, 3)],
        )

    monkeypatch.setattr(document, "is_grobid_available", lambda: True)
    monkeypatch.setattr(document, "extract_document_via_grobid", fake_grobid)
    monkeypatch.setattr(document, "_read_source_text", lambda path: _SOURCE)

    extracted = extract_document(_pdf(tmp_path))

    assert len(calls) == 1
    assert [entry.title for entry in extracted.references] == ["A Title"]
    assert [context.marker_text for context in extracted.citations] == ["[1]"]


def test_a_bibliography_only_call_never_reads_the_body(monkeypatch, tmp_path):
    """`extract_references` stays the cheaper entry point: on the GROBID path it must not
    touch pymupdf4llm at all, which is the behaviour every existing caller has today."""
    monkeypatch.setattr(document, "is_grobid_available", lambda: True)
    monkeypatch.setattr(
        document,
        "extract_document_via_grobid",
        lambda path: GrobidDocument(references=[RawReferenceEntry(1, "Doe, J. (2020) A Title")]),
    )
    monkeypatch.setattr(
        document, "_read_source_text", lambda path: pytest.fail("the body was read for a bibliography-only call")
    )

    assert len(extract_references(_pdf(tmp_path))) == 1


def test_falling_back_to_tier_0_drops_grobids_citations_with_its_references(monkeypatch, tmp_path):
    """A reference list that came from Tier 0 has no TEI ids in it, so GROBID's markers
    point into a list nobody is holding. Carrying them forward would hand the matcher
    targets that resolve against the wrong document."""
    monkeypatch.setattr(document, "is_grobid_available", lambda: True)
    monkeypatch.setattr(
        document,
        "extract_document_via_grobid",
        lambda path: GrobidDocument(citations=[CitationContext(1, "b0", "[1]", "a paragraph", 0, 3)]),
    )
    monkeypatch.setattr(document, "_read_source_text", lambda path: _SOURCE)

    extracted = extract_document(_pdf(tmp_path))

    assert extracted.citations == ()
    assert len(extracted.references) == 2


def test_a_grobid_failure_falls_back_with_the_body_intact(monkeypatch, tmp_path):
    def raise_unavailable(path):
        raise GrobidUnavailableError("connection refused")

    monkeypatch.setattr(document, "is_grobid_available", lambda: True)
    monkeypatch.setattr(document, "extract_document_via_grobid", raise_unavailable)
    monkeypatch.setattr(document, "_read_source_text", lambda path: _SOURCE)

    extracted = extract_document(_pdf(tmp_path))

    assert len(extracted.references) == 2
    assert extracted.citations == ()
    assert "Body text citing [1] here." in extracted.artifact.body_text


def test_engine_anchor_skips_grobid_here_too(monkeypatch, tmp_path):
    def _fail_if_called(*_args, **_kwargs):
        raise AssertionError("GROBID must not be touched when ENGINE_ANCHOR is forced")

    monkeypatch.setattr(document, "is_grobid_available", _fail_if_called)
    monkeypatch.setattr(document, "extract_document_via_grobid", _fail_if_called)
    monkeypatch.setattr(document, "_read_source_text", lambda path: _SOURCE)

    assert len(extract_document(_pdf(tmp_path), engine=ENGINE_ANCHOR).references) == 2


def test_engine_grobid_skips_the_reachability_check_here_too(monkeypatch, tmp_path):
    def _fail_if_called():
        raise AssertionError("is_grobid_available should not be consulted when ENGINE_GROBID is forced")

    monkeypatch.setattr(document, "is_grobid_available", _fail_if_called)
    monkeypatch.setattr(
        document,
        "extract_document_via_grobid",
        lambda path: GrobidDocument(references=[RawReferenceEntry(1, "Doe, J. (2020) A Title")]),
    )
    monkeypatch.setattr(document, "_read_source_text", lambda path: _SOURCE)

    assert len(extract_document(_pdf(tmp_path), engine=ENGINE_GROBID).references) == 1


# --------------------------------------------------------------------------------------
# No bibliography heading
# --------------------------------------------------------------------------------------


def test_a_document_with_no_bibliography_heading_still_fails_for_tier_0(monkeypatch, tmp_path):
    """Unchanged from extract_references: without a reference list there is nothing to
    check a citation against, and the caller is entitled to hear so."""
    monkeypatch.setattr(document, "_read_source_text", lambda path: "Body text with no reference list.")

    with pytest.raises(NoBibliographySectionError):
        extract_document(_pdf(tmp_path))


def test_a_document_with_no_bibliography_heading_still_yields_a_body_when_grobid_read_it(
    monkeypatch, tmp_path
):
    """A manuscript whose reference list this project cannot locate is exactly the one
    whose in-text citations are worth showing — so a missing heading is only fatal when
    the references depended on it."""
    monkeypatch.setattr(document, "is_grobid_available", lambda: True)
    monkeypatch.setattr(
        document,
        "extract_document_via_grobid",
        lambda path: GrobidDocument(references=[RawReferenceEntry(1, "Doe, J. (2020) A Title")]),
    )
    monkeypatch.setattr(document, "_read_source_text", lambda path: "Body text citing (Doe, 2020) with no list.")

    extracted = extract_document(_pdf(tmp_path))

    assert len(extracted.references) == 1
    assert "(Doe, 2020)" in extracted.artifact.body_text
    assert extracted.profile is None


# --------------------------------------------------------------------------------------
# What the Phase 4 review pass found (see the plan's Notes)
# --------------------------------------------------------------------------------------


def test_a_body_the_parser_cannot_read_does_not_take_the_reference_check_with_it(monkeypatch, tmp_path):
    """The plan's first Success Criterion says the reference output holds "when citation
    detection/matching fails **or its inputs are unavailable**". The artifact is an input.

    Reading the body is new work this phase added to a path that did not do it: a PDF
    GROBID can parse but pymupdf4llm cannot open — encrypted, truncated — used to produce
    a complete reference check, and briefly produced a failed one instead.
    """
    monkeypatch.setattr(document, "is_grobid_available", lambda: True)
    monkeypatch.setattr(
        document,
        "extract_document_via_grobid",
        lambda path: GrobidDocument(references=[RawReferenceEntry(1, "Doe, J. (2020) A Title", title="A Title")]),
    )

    def unreadable(path):
        raise RuntimeError("cannot open manuscript.pdf as pdf")

    monkeypatch.setattr(document, "_extract_pdf_markdown", unreadable)
    path = _pdf(tmp_path)

    extracted = extract_document(path)

    assert extracted.references == extract_references(path)
    assert extracted.artifact is None and extracted.profile is None


def test_a_body_that_cannot_be_turned_into_an_artifact_is_survivable_too(monkeypatch, tmp_path):
    """The same rule one step later: by the time the references exist, nothing the body
    does may endanger them."""
    monkeypatch.setattr(document, "_read_source_text", lambda path: _SOURCE)

    def broken(raw_text, suffix):
        raise ValueError("artifact construction fell over")

    monkeypatch.setattr(document, "_artifact_from", broken)

    extracted = extract_document(_pdf(tmp_path))

    assert len(extracted.references) == 2
    assert extracted.artifact is None


def test_an_unreadable_document_still_fails_when_the_references_depend_on_it(monkeypatch, tmp_path):
    """The guard must not swallow the Tier 0 case: with no GROBID result there is nothing
    to protect, and a caller that gets no references is entitled to hear why."""
    def unreadable(path):
        raise RuntimeError("cannot open manuscript.pdf as pdf")

    monkeypatch.setattr(document, "_extract_pdf_markdown", unreadable)

    with pytest.raises(RuntimeError):
        extract_document(_pdf(tmp_path))

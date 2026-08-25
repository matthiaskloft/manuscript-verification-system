"""Locate a document's bibliography section and hand it to tier0 splitting.

The missing link between a PDF/DOCX file and refcheck.extraction.tier0, which only
splits bibliography text that's already been isolated from the rest of the document.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from refcheck.extraction.document_artifact import (
    DocumentArtifact,
    SourceLine,
    concat_tag,
    retag,
    segments_of,
)
from refcheck.extraction.engine_status import ENGINE_ANCHOR, ENGINE_AUTO, ENGINE_GROBID
from refcheck.extraction.grobid import (
    CitationContext,
    GrobidUnavailableError,
    extract_document_via_grobid,
    is_grobid_available,
)
from refcheck.extraction.reference_list_audit import (
    NOT_AUDITED,
    ReferenceListAudit,
    audit_reference_list,
)
from refcheck.extraction.signals import (
    APA_AUTHOR_START_RE,
    BARE_NUMBER_RE,
    BARE_YEAR_RE,
    BRACKET_NUMBER_RE,
    DANGLING_HYPHEN_RE,
    DATE_SLOT_RE,
    DOT_NUMBER_RE,
    JSS_AUTHOR_START_RE,
    LINK_END_RE,
    LOCATOR_CONTINUATION_RE,
    entry_is_complete,
)
from refcheck.extraction.style_profile import StyleProfile, compute_style_profile
from refcheck.extraction.tier0 import RawReferenceEntry, split_bibliography_block

# pymupdf4llm renders each page's footer page-number as its own Markdown paragraph —
# a bare number with a blank line before *and* after it. Requiring blank lines on
# *both* sides (not just "looks like a short number") keeps this from ever touching a
# genuine bare-number line that's part of a real entry's own wrapped text (a page
# range, volume number, or a year wrapped onto its own line), since those don't sit
# alone in their own blank-delimited paragraph the way a footer does.
_PAGE_NUMBER_LINE_RE = re.compile(r"^\s*\d{1,4}\s*$")

# A bibliography entry can cross a physical page boundary in the middle of a long
# author list. pymupdf4llm then emits the page number as an isolated paragraph, with
# a blank paragraph on either side. The continuation can itself look like a new APA
# entry because it starts with another author's surname, so the ordinary entry-start
# heuristic cannot distinguish it. A trailing separator is the stronger signal: a
# real completed reference should not end in a comma/semicolon/colon or dangling
# hyphen immediately before the page number.
_PAGE_BREAK_CONTINUATION_RE = re.compile(r"[,;:\-–—]\s*$")

# A running header/footer that isn't a bare number (a journal name + article title,
# e.g. "www.annualreviews.org * Sample Size Planning 563") repeats on every page with
# only its page number differing, so _PAGE_NUMBER_LINE_RE alone misses it — observed
# on a real Annual Review article where it survived as extra spurious "entries", one
# per page, at the end of the reference list. Collapsing digit runs before comparing
# turns each page's copy into the same key, so >=2 isolated (blank-delimited) matches
# anywhere in the document is treated as a running header/footer and dropped. Two
# distinct real one-line bibliography entries coincidentally sharing identical
# non-digit text is not a realistic risk; bulleted entries are excluded from this
# check entirely since a Markdown list item is never a running header/footer.
_DIGIT_RUN_RE = re.compile(r"\d+")
_MIN_REPEATS_TO_TREAT_AS_RUNNING_LINE = 2

# ...but two occurrences is a low bar, and body text clears it constantly: a table note
# repeated under two tables, a methods sentence repeated across two studies, two figure
# captions differing only in their number. All were being deleted outright — four
# ordinary paragraphs, two of them a repeated sentence, normalized to "". That fails in
# the worst available direction, because a body line lost here becomes a reference
# nothing cites, i.e. a reviewer-facing "unused reference" that isn't one.
#
# What separates the two is shape, not count: a running header/footer is a label, and a
# label is not a sentence. Requiring terminal punctuation *and* several words keeps the
# motivating case droppable ("www.annualreviews.org * Sample Size Planning 563" ends in
# a page number, and a bare repeated heading has no terminal punctuation either) while
# putting prose out of reach. The residual cost is accepted deliberately: a footer that
# does read like a sentence ("Downloaded from example.org on 12 January 2020.") now
# survives into the reference list, where it may become one spurious entry. One
# spurious entry is a visible, checkable row; silently deleted body text is neither.
_PROSE_TERMINAL_RE = re.compile(r"[.!?][\"'’”)\]]*$")
_MIN_PROSE_WORDS = 4

# pymupdf4llm renders headings as "#"/"##"/... Markdown, and often wraps stray text in
# "**bold**" (its font-weight-detection heuristic also fires on non-heading bold body
# text). Bibliography entries come out as Markdown list items ("- [1] ..."). None of
# that is known to _is_heading's exact-line match or to tier0's marker regexes
# (extraction/tier0.py), which both predate Markdown output and expect a plain
# "[1] ..." or "References" line — so all three markers are stripped back to plain
# text before that unchanged logic runs.
_MD_HEADING_PREFIX_RE = re.compile(r"^#{1,6}\s*")
_MD_LIST_BULLET_RE = re.compile(r"^\s*[-*]\s+")
_MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_MD_BLOCKQUOTE_PREFIX_RE = re.compile(r"^\s*>\s*")
_MD_TABLE_NUMBER_MARKER_RE = re.compile(r"\*\*(\d{1,3})[.)]\*\*\s*(?:<br\s*/?>)?", re.IGNORECASE)

# pymupdf4llm also renders italics as "_text_" (its font-detection heuristic fires on
# journal/book titles, which are usually italicized in APA-style references) and
# occasionally mis-renders an unusual glyph run as Markdown strikethrough ("~~...~~")
# or a literal HTML tag (e.g. "<sup>...</sup>" for a raised diacritic) — none of which
# are meaningful here (unlike GROBID, which never shows the original PDF text at all;
# see grobid.py's _format_raw_text, which reconstructs a short "Author (Year) Title"
# string instead). Left unstripped, these show up verbatim in the UI's raw-citation
# column, and this tier's raw_text can end up considerably longer and messier than
# GROBID's per entry as a result.
_MD_ITALIC_RE = re.compile(r"_(.+?)_")
_MD_STRIKETHROUGH_RE = re.compile(r"~~(.+?)~~")
_HTML_TAG_RE = re.compile(r"</?[a-zA-Z][^<>]*>")

SUPPORTED_SUFFIXES = {".pdf", ".docx"}

# Headings that mark the start of a bibliography, matched against a stripped,
# otherwise-standalone line (how a heading typically renders once PDF/DOCX text is
# flattened to plain text). English + German, since the project targets a German
# university deployment (see docs/project-plan.md).
_SECTION_HEADINGS = {
    "references",
    "reference list",
    "bibliography",
    "works cited",
    "literature cited",
    "literaturverzeichnis",
    "literatur",
    "quellenverzeichnis",
}

# Headings that, if found after the bibliography heading, mark where it ends — so an
# appendix or acknowledgements section doesn't get swept into the reference list.
_STOP_HEADINGS = {
    "appendix",
    "appendices",
    "acknowledgements",
    "acknowledgments",
    "supplementary material",
    "author",
    "authors",
    "author note",
    "author notes",
    "author information",
    "anhang",
    "danksagung",
}


# A lettered/numbered appendix heading ("Appendix A.", "Appendix B Additional
# analyses") and the submission-history metadata line journals print after the
# reference list ("Manuscript received: March 5, 2020"). Both end the bibliography,
# but both also have to stay heading-shaped: an unbounded trailing ".+" would let a
# *reference* that happens to open with "Appendix"/"Received" truncate everything
# after it. Requiring a short tail with no sentence-internal period keeps a real
# entry — which carries author, title, and journal, each period-separated — out.
_HEADING_TAIL = r"[^.]{0,60}"
# The appendix designator stays case-sensitive ("Appendix B", "Appendix IV"), and a
# descriptive tail is only allowed after a designator or a colon/period. Without that,
# a lowercase word would qualify and a wrapped book title ("Appendix to the Hearings,
# Volume XXVI") would read as a heading. "received" likewise needs a date to look like
# submission history rather than the opening words of a title.
_NAMED_APPENDIX_RE = re.compile(
    rf"[Aa]ppendi(?:x|ces)(?:[.:]\s*{_HEADING_TAIL}|\s+[A-Z0-9IVX]{{1,4}}\b[.:]?\s*{_HEADING_TAIL})\.?"
)
_RECEIVED_METADATA_RE = re.compile(
    rf"(?:manuscript\s+)?received\b{_HEADING_TAIL}(?:1[5-9]\d{{2}}|20\d{{2}})[^.]{{0,20}}\.?", re.IGNORECASE
)


class NoBibliographySectionError(ValueError):
    """Raised when no recognizable bibliography heading is found in the document."""


def extract_full_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix == ".docx":
        return _extract_docx_text(path)
    raise ValueError(f"Unsupported file type: {suffix!r} (expected .pdf or .docx)")


def _extract_pdf_text(path: Path) -> str:
    return normalize_markdown_text(_extract_pdf_markdown(path))


def _extract_pdf_markdown(path: Path) -> str:
    """Uses pymupdf4llm's Markdown extraction rather than raw fitz get_text(): real
    Markdown headings (derived from font size) make bibliography-heading detection
    more robust than plain-text heuristics, and pymupdf4llm's explicit multi-column
    layout handling avoids the interleaved-text failure mode plain get_text() has on
    two-column manuscripts. See docs/project-plan.md's pymupdf4llm TODO — this replaced the old
    fitz-based implementation on the strength of that comparison.

    Returned before normalization so build_document_artifact can keep the original as
    the other end of its source-anchor mapping. _extract_pdf_text normalizes exactly
    this string, so the two paths cannot extract differently.
    """
    import pymupdf4llm

    return pymupdf4llm.to_markdown(str(path))


def normalize_markdown_text(md: str) -> str:
    """The Markdown cleanup chain, as a function of text rather than of a PDF path.

    Separated from _extract_pdf_text so these passes can be exercised from a text
    fixture. They are the earliest layer that can lose a whole reference — a bad merge
    here deletes an entry before any splitting heuristic ever sees it — and until this
    split they could only be reached by compiling a PDF that reproduced the artifact.
    """
    return "\n".join(_normalize_markdown_lines(md.splitlines()))


def _normalize_markdown_lines(raw_lines: list[str]) -> list[SourceLine]:
    """normalize_markdown_text's chain, keeping each output line's raw-line index.

    The passes below drop, merge and rewrite lines, so an output line's position says
    nothing about where it came from. Tagging the input once here and letting
    SourceLine ride through (see document_artifact.SourceLine) is what lets
    build_document_artifact anchor a normalized span back to the original text without
    any pass having to maintain a second, parallel index structure.
    """
    lines: list[str] = [SourceLine(line, i) for i, line in enumerate(raw_lines)]
    lines = _expand_markdown_table_reference_rows(lines)
    lines = _merge_page_break_continuations(lines)
    lines = _drop_isolated_page_numbers(lines)
    lines = _drop_isolated_decorative_list_markers(lines)
    lines = _drop_repeated_isolated_lines(lines)
    return _merge_mid_entry_paragraph_breaks(lines)


def _expand_markdown_table_reference_rows(lines: list[str]) -> list[str]:
    """Turn numbered bibliography entries embedded in Markdown table cells into lines."""
    expanded: list[str] = []
    for line in lines:
        if "|" not in line or not _MD_TABLE_NUMBER_MARKER_RE.search(line):
            expanded.append(line)
            continue

        flattened = line.strip().strip("|").replace("|", " ")
        flattened = re.sub(r"<br\s*/?>", " ", flattened, flags=re.IGNORECASE)
        flattened = _MD_TABLE_NUMBER_MARKER_RE.sub(lambda match: f"\n{match.group(1)}. ", flattened)
        expanded.extend(retag(part.strip(), line) for part in flattened.splitlines() if part.strip())
    return expanded


def _merge_page_break_continuations(lines: list[str]) -> list[str]:
    """Join text split around an isolated page number when the preceding fragment
    ends with continuation punctuation.

    This deliberately does not merge every unfinished-looking line across a page:
    references ending in a bare DOI sometimes omit the final period, and blindly
    joining those would absorb the next real reference. The punctuation requirement
    targets the observed long-author-list failure while preserving that boundary.
    """
    result = lines.copy()
    for i, line in enumerate(lines):
        if not (_PAGE_NUMBER_LINE_RE.match(line) and _is_isolated(lines, i)):
            continue

        previous = next((j for j in range(i - 1, -1, -1) if result[j].strip()), None)
        following = next((j for j in range(i + 1, len(lines)) if result[j].strip()), None)
        if previous is None or following is None:
            continue
        if not _PAGE_BREAK_CONTINUATION_RE.search(result[previous]):
            continue

        # The merged line spans two pages, so it has two origins: concat_tag keeps the
        # continuation's own, and a span landing in it resolves to the page it is on.
        head = result[previous].rstrip()
        tail = result[following]
        result[previous] = concat_tag(
            f"{head} {tail.lstrip()}", result[previous], tail, len(head) + 1, len(tail) - len(tail.lstrip())
        )
        result[following] = retag("", result[following])
    return result


def _is_isolated(lines: list[str], i: int) -> bool:
    def is_blank(j: int) -> bool:
        return j < 0 or j >= len(lines) or not lines[j].strip()

    return bool(lines[i].strip()) and is_blank(i - 1) and is_blank(i + 1)


def _drop_isolated_page_numbers(lines: list[str]) -> list[str]:
    return [line for i, line in enumerate(lines) if not (_PAGE_NUMBER_LINE_RE.match(line) and _is_isolated(lines, i))]


def _drop_isolated_decorative_list_markers(lines: list[str]) -> list[str]:
    """Drop text-free Markdown list artifacts in otherwise empty paragraphs.

    A standalone bullet marker cannot be a complete bibliography entry. Keeping it
    causes tier0 to emit a spurious reference when extraction represents a PDF
    decoration as ``- x`` (a single leftover glyph) or as an empty ``- `` bullet.
    Both carry at most one character once the marker itself is stripped, so both
    are treated the same. The isolation requirement (a blank line on either side,
    as in _drop_isolated_page_numbers) keeps a bullet that is part of a real list
    out of reach even if its own text is that short.
    """

    def is_decorative(line: str) -> bool:
        return bool(_MD_LIST_BULLET_RE.match(line)) and len(_normalize_markdown_line(line).strip()) <= 1

    return [line for i, line in enumerate(lines) if not (is_decorative(line) and _is_isolated(lines, i))]


def _drop_repeated_isolated_lines(lines: list[str]) -> list[str]:
    """See _DIGIT_RUN_RE's comment for the running-header/footer case this targets, and
    _PROSE_TERMINAL_RE's for why a sentence is exempt from it however often it repeats.

    A prose-shaped line is excluded from the candidate set entirely, not merely spared
    at the end: it must not count toward another line's repeat total either, or two
    table notes would still conspire to delete a third thing that happens to share
    their key.

    A line that itself looks like a real bibliography entry's own start (matches one
    of tier0's numbered/APA/JSS markers) is never eligible to be dropped here, even if
    it repeats — a genuinely short, formulaic entry (e.g. several editions of the same
    numbered standard, "[3] ISO 8601:2019...", "[7] ISO 8601:2022...") could otherwise
    canonicalize to an identical key and coincidentally clear the repeat threshold,
    which would be a real bibliography entry lost, not a footer removed.

    A repeated section or stop heading is handled by keeping its *first* occurrence
    and dropping the rest, which is not the same as exempting it. A running header
    repeating "References" on every page made every copy of that word — including the
    real heading — clear the threshold, and find_bibliography_section then found no
    bibliography at all: its own defense against that header (take the first
    occurrence) never got the chance to run, because this pass had already deleted the
    line it would have matched. Exempting headings outright instead would keep the
    real one but leave every running-header copy sitting inside the reference list,
    where tier0 turns each into a spurious entry. First-wins is that defense, stated
    where the deletion actually happens.

    Every check here canonicalizes through _normalize_markdown_line. It did not
    always: `canonical` used the raw line while `looks_like_entry_start` normalized,
    so a "## References" heading and a plain "References" header hashed differently
    and the collision above was masked for Markdown-headed documents but not for
    plain-text ones. Two normalization disciplines inside one pass is how that stayed
    hidden.
    """

    def canonical(line: str) -> str:
        return _DIGIT_RUN_RE.sub("#", _normalize_markdown_line(line).strip().casefold())

    def looks_like_entry_start(line: str) -> bool:
        normalized = _normalize_markdown_line(line)
        return any(
            marker_re.match(normalized)
            for marker_re in (BRACKET_NUMBER_RE, DOT_NUMBER_RE, BARE_NUMBER_RE, APA_AUTHOR_START_RE, JSS_AUTHOR_START_RE)
        )

    def reads_like_prose(line: str) -> bool:
        text = _normalize_markdown_line(line).strip()
        return bool(_PROSE_TERMINAL_RE.search(text)) and len(text.split()) >= _MIN_PROSE_WORDS

    def is_section_heading(line: str) -> bool:
        text = _normalize_markdown_line(line).strip()
        return bool(text) and (_is_heading(text, _SECTION_HEADINGS) or _is_heading(text, _STOP_HEADINGS))

    isolated_idx = [
        i
        for i in range(len(lines))
        if _is_isolated(lines, i)
        and not _MD_LIST_BULLET_RE.match(lines[i])
        and not looks_like_entry_start(lines[i])
        and not reads_like_prose(lines[i])
    ]
    counts: dict[str, int] = {}
    first_seen: dict[str, int] = {}
    for i in isolated_idx:
        key = canonical(lines[i])
        counts[key] = counts.get(key, 0) + 1
        first_seen.setdefault(key, i)

    def is_the_real_heading(i: int) -> bool:
        return is_section_heading(lines[i]) and first_seen[canonical(lines[i])] == i

    drop = {
        i
        for i in isolated_idx
        if counts[canonical(lines[i])] >= _MIN_REPEATS_TO_TREAT_AS_RUNNING_LINE
        and not is_the_real_heading(i)
    }
    return [line for i, line in enumerate(lines) if i not in drop]


def _merge_mid_entry_paragraph_breaks(lines: list[str]) -> list[str]:
    """Undo a pymupdf4llm artifact: for a plain-paragraph bibliography (no numbering
    or Markdown bullets — an APA/JSS-style running-text list), pymupdf4llm's block
    segmentation can insert a blank line in the middle of a single wrapped entry at a
    column or page boundary, not just between entries (observed on a real two-column
    journal article: "...requirements for accurate estimation of\n\nsquared
    semipartial correlation coefficients..." — one entry, split mid-noun-phrase).

    tier0's blank-line-separated-entries heuristic can't tell that apart from a real
    entry boundary and, worse, runs *before* the more specific APA/JSS author-start
    heuristics in split_bibliography_block — so left uncorrected, spurious blank
    lines don't just add a stray entry, they hijack the whole split away from the
    heuristic that would have gotten it right.

    Markdown list items ("- ...", one per bibliography entry when pymupdf4llm detects
    a numbered/bulleted layout) are usually already one complete entry each — but a
    page break falling mid-entry can make pymupdf4llm emit the continuation as its own
    separate bullet too (observed on a real PDF: a hyphenated word wrapped across a
    page, and separately a title-ending period immediately followed by a page break
    before the journal name — both landed in their own bullet). Two bullets separated
    by a blank line are therefore eligible for the same merge as plain paragraphs, via
    the `has_own_evidence` check below, which is stricter than the plain-text
    marker check: a bulleted list's real entries reliably open with an author-start
    marker *or* carry their own parenthesised year (or "n.d."/"in press" equivalent),
    so requiring one or the other of those before treating a bullet as a new entry
    catches the corporate-author case the marker regexes alone miss (e.g. "World
    Health Organization. (2021)..."), where a plain-paragraph bibliography instead
    relies on entry_is_complete of the *previous* line to avoid over-merging — a signal
    that doesn't work for bullets, since an ordinary title-ending period before the
    journal name (i.e., mid-entry) is indistinguishable from a genuinely finished
    entry by punctuation alone.

    Lacking its own marker/year is not sufficient on its own, though — a genuinely new
    but *undated* entry (a corporate author or title-first reference with no year at
    all) looks exactly like a continuation fragment by that measure alone, and would
    otherwise be silently swallowed into whichever bullet came before it. Merging a
    bulleted pair therefore also requires positive continuation evidence: the previous
    fragment ends in a dangling hyphen (a word wrapped across the break) or the current
    one carries a journal-locator shape ("34(2), 1-22") or a DOI/URL — the two forms a
    continuation fragment actually takes, as opposed to a new entry's own prose.
    A blank line is otherwise only ever collapsed here when *neither* surrounding line
    is a recognized section heading.

    Two independent signals decide whether the remaining, non-bulleted blank lines
    are genuine boundaries: the line *after* the blank must not look like a new
    entry's own start (tier0's own numbered/APA/JSS marker regexes, including the
    bare-number one — reused here, not reimplemented, so both places agree on what a
    marker looks like), and the line *before* it must already look finished —
    sentence-terminal punctuation *or* a trailing DOI/URL, the two ways an entry
    legitimately ends, the same pair _is_stop_heading uses.

    The link-end half of that was missing until a DOI-final entry was found being
    merged into the corporate-author entry after it, deleting a whole reference before
    splitting ever saw it. The marker check was believed to cover the DOI case, and
    does whenever the *next* entry starts with a recognized marker ("2. ", or just "2 "
    in a bare-number bibliography) — but "World Health Organization. (2021)..." matches
    none of them, and neither does any other corporate author. Omitting the
    bare-number marker here (unlike tier0's own use of it, which additionally requires
    strictly sequential numbers before trusting it) would silently corrupt a
    bare-number bibliography: a false merge across two real entries breaks the
    1, 2, 3, ... sequence split_bibliography_block requires, defeating that heuristic
    entirely for the rest of the document.
    """
    normalized = [_normalize_tagged_line(line) for line in lines]
    is_bullet = [bool(_MD_LIST_BULLET_RE.match(line)) for line in lines]
    is_heading = [
        bool(text) and (_is_heading(text, _SECTION_HEADINGS) or _is_heading(text, _STOP_HEADINGS))
        for text in normalized
    ]

    merged: list[str] = []
    merged_is_bullet: list[bool] = []
    merged_is_heading: list[bool] = []
    pending_blank = False

    for text, bullet, heading in zip(normalized, is_bullet, is_heading, strict=True):
        if not text.strip():
            pending_blank = pending_blank or bool(merged)
            continue
        if pending_blank:
            looks_like_new_entry = any(
                marker_re.match(text)
                for marker_re in (
                    BRACKET_NUMBER_RE,
                    DOT_NUMBER_RE,
                    BARE_NUMBER_RE,
                    APA_AUTHOR_START_RE,
                    JSS_AUTHOR_START_RE,
                )
            )
            if merged and merged_is_bullet[-1] and bullet:
                # Both sides are bullets: entry_is_complete on the previous fragment
                # can't discriminate (see the docstring), so a bulleted continuation is
                # recognized by the *absence* of its own entry-start evidence instead.
                # A style without a comma-initial or JSS-style marker (e.g. Chicago
                # author-date's "Surname, Full Name. YYYY.") still opens with a bare
                # year right after the author, which a genuine continuation fragment
                # (a bare journal name, a page range, a trailing DOI) never carries —
                # checking for either kind of year, not just the parenthesised one
                # DATE_SLOT_RE covers, is what keeps this from merging that style's own
                # distinct entries together. The bare-year check is restricted to the
                # opening of the line — searching the whole fragment would also catch a
                # year incidental to a continuation (a large page range, a cited year
                # inside a title), which would wrongly block the merge this is meant to
                # allow.
                leading = text.lstrip()[:40]
                has_own_evidence = looks_like_new_entry or DATE_SLOT_RE.search(text) or BARE_YEAR_RE.search(leading)
                has_continuation_evidence = bool(
                    DANGLING_HYPHEN_RE.search(merged[-1])
                    or LOCATOR_CONTINUATION_RE.search(text)
                    or LINK_END_RE.search(text)
                )
                can_merge = (
                    not merged_is_heading[-1]
                    and not heading
                    and not has_own_evidence
                    and has_continuation_evidence
                )
            else:
                can_merge = (
                    merged
                    and not merged_is_bullet[-1]
                    and not merged_is_heading[-1]
                    and not bullet
                    and not heading
                    and not looks_like_new_entry
                    and not entry_is_complete(merged[-1])
                )
            if can_merge:
                # Over body text this is the pass that folds a heading into the prose
                # under it, so the fragment's own origin is the one nearly every span
                # in the result wants; concat_tag is what keeps it.
                head = merged[-1].rstrip()
                merged[-1] = concat_tag(
                    f"{head} {text.strip()}", merged[-1], text, len(head) + 1, len(text) - len(text.lstrip())
                )
                pending_blank = False
                continue
            merged.append(retag("", text))
            merged_is_bullet.append(False)
            merged_is_heading.append(False)
            pending_blank = False
        merged.append(text)
        merged_is_bullet.append(bullet)
        merged_is_heading.append(heading)

    return merged


def _strip_markdown_markup(line: str) -> str:
    """_normalize_markdown_line without the closing rstrip.

    Split out so _normalize_tagged_line can ask how long a *prefix* of a line becomes
    once markup is removed, which is how a segment boundary crosses this rewrite. The
    rstrip has to stay out of that question: it would eat the boundary's own trailing
    space and pull the following fragment's offset back into the fragment before it.
    """
    stripped = _MD_BLOCKQUOTE_PREFIX_RE.sub("", line)
    stripped = _MD_HEADING_PREFIX_RE.sub("", stripped)
    stripped = _MD_LIST_BULLET_RE.sub("", stripped)
    stripped = _MD_BOLD_RE.sub(r"\1", stripped)
    stripped = _MD_ITALIC_RE.sub(r"\1", stripped)
    stripped = _MD_STRIKETHROUGH_RE.sub(r"\1", stripped)
    return _HTML_TAG_RE.sub("", stripped)


def _normalize_markdown_line(line: str) -> str:
    return _strip_markdown_markup(line).rstrip()


def _normalize_tagged_line(line: str) -> SourceLine:
    """_normalize_markdown_line, carrying the line's segment structure through it.

    Markup removal shortens the text, so a segment boundary at raw offset *k* lands
    wherever the markup-stripped prefix ends. Boundaries only ever move left, and a
    line that came from a single raw line — the overwhelming majority — skips the
    mapping entirely.
    """
    text = _normalize_markdown_line(line)
    segments = segments_of(line)
    if len(segments) == 1:
        return retag(text, line)
    rebased = [
        (min(len(_strip_markdown_markup(line[:offset])), len(text)), origin) for offset, origin in segments
    ]
    return SourceLine(text, None, tuple(rebased))


def _extract_docx_text(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _is_heading(line: str, headings: set[str]) -> bool:
    return line.strip().strip(":").lower() in headings


def _is_stop_heading(line: str, previous: str) -> bool:
    """Whether `line` ends the bibliography, given the last non-blank line before it.

    The patterns below match a line's whole text, but a reference wraps across lines,
    and any one of those lines can look heading-shaped on its own ("Received signal
    strength fingerprints" from a wireless-positioning title, "Appendix to the
    Hearings, Volume XXVI" from a book title). A heading only ever follows a
    *completed* entry, so a predecessor that doesn't end in sentence-terminal
    punctuation means this line is a continuation, not a boundary — mid-entry is
    exactly where a false stop would silently discard the rest of the bibliography.

    An entry ending in a bare DOI or URL counts as completed too: APA closes those
    without a period, and the last entry before an appendix is as likely to end that
    way as any other, so punctuation alone would miss the real heading after it.
    """
    normalized = line.strip().strip(":")
    if _is_heading(normalized, _STOP_HEADINGS):
        return True
    if previous and not entry_is_complete(previous):
        return False
    return bool(_NAMED_APPENDIX_RE.fullmatch(normalized) or _RECEIVED_METADATA_RE.fullmatch(normalized))


def find_bibliography_section(full_text: str) -> str:
    """Return the bibliography text: everything between the first recognized heading
    and the next stop heading (or end of document).

    Uses the *first* matching heading rather than the last: a running header repeating
    "References" on every page of a multi-page bibliography would otherwise cause the
    last-occurrence approach to skip every entry before that final repeat.

    Known limitation: a heading is matched by exact string equality against a whole,
    stripped line — deliberately strict, to avoid matching a table-of-contents entry
    like "References ..... 12". This cuts both ways: a manuscript whose body text
    happens to wrap so that a line reads exactly "References" would false-positive
    (e.g. "...documented in the\nReferences\nsection above."), while a heading that
    itself wraps across two lines during PDF text flattening (e.g. "Reference\nList")
    would false-negative and raise NoBibliographySectionError even though a real
    bibliography section exists — arguably the more likely failure in practice. Not
    seen on the synthetic fixtures this was tested against; no defense against either
    case exists yet.
    """
    lines = full_text.splitlines()
    bounds = _bibliography_bounds(lines)
    if bounds is None:
        raise NoBibliographySectionError("No bibliography/references heading found in the document.")
    heading_idx, end_idx = bounds
    return "\n".join(lines[heading_idx + 1 : end_idx]).strip()


def _bibliography_bounds(lines: list[str]) -> tuple[int, int] | None:
    """The bibliography's (heading index, end index) in `lines`, or None if absent.

    Split out of find_bibliography_section so build_document_artifact can ask the
    complementary question — which lines are *not* the bibliography — off the same
    answer. Two independent notions of where the reference list stops is precisely the
    divergence signals.py's docstring is about.
    """
    heading_idx = next((i for i, line in enumerate(lines) if _is_heading(line, _SECTION_HEADINGS)), None)
    if heading_idx is None:
        return None

    # Carried forward rather than re-scanned backwards per line: the backward search
    # is quadratic on a section holding a long run of blank lines.
    end_idx = len(lines)
    previous = ""
    for i in range(heading_idx + 1, len(lines)):
        if _is_stop_heading(lines[i], previous):
            end_idx = i
            break
        if lines[i].strip():
            previous = lines[i].strip()
    return heading_idx, end_idx


@dataclass(frozen=True)
class ExtractedDocument:
    """One document, read once: its bibliography and everything citation matching needs.

    `citations` is Tier 1's output. It is empty whenever GROBID did not produce the
    reference list — but also when GROBID read the document and found no `<ref
    type="bibr">` markers in it at all, which a scanned PDF or a body its segmentation
    model missed will do. So "empty" means "no Tier 1 markers to work from", not "GROBID
    was not here, and `citation_matching` falling back to Tier 0 on it is a deliberate
    second chance rather than a statement about the extractor. That fallback matches
    page-scraped markers against GROBID's *reconstructed* entry text, which carries no
    printed reference numbers — so a numbered marker resolves through position, and a
    GROBID reference list shorter than the printed one resolves it to the wrong entry.

    `artifact` and `profile` are None when the caller did not ask for the body (see
    `extract_references`), when the document has no recognizable bibliography heading, or
    when reading the body failed on a document whose references came from GROBID anyway.

    `audit` is how `references` compares to the list the document printed. It is unread
    (`ReferenceListAudit.read` False) in exactly those same cases, since re-reading the
    printed list needs the bibliography text and that is only extracted for the body — so
    a caller on the cheap path is told "not checked" rather than "checked and fine".
    """

    references: list[RawReferenceEntry]
    citations: tuple[CitationContext, ...] = ()
    artifact: DocumentArtifact | None = None
    profile: StyleProfile | None = None
    audit: ReferenceListAudit = NOT_AUDITED


def extract_document(path: Path, engine: str = ENGINE_AUTO) -> ExtractedDocument:
    """Read a document once for everything a check needs: references, body, citations.

    The same extraction `extract_references` does, keeping what it discards. Worth having
    as its own entry point rather than as a flag on that one because the two differ in
    cost, not just in output: this reads the PDF with pymupdf4llm even when GROBID
    supplied the references, since the TEI GROBID returns shares no coordinates with a
    `DocumentArtifact` and a marker has to be findable in the manuscript's own text.

    What it costs, stated precisely because the first version of this docstring overstated
    it. Before this existed, a caller wanting both halves called `extract_references` and
    then `build_document_artifact`: one GROBID round-trip either way (Phase 3 already made
    `extract_references_via_grobid` a view onto one call), plus two PDF parses on the
    Tier 0 path and one on the GROBID path. This is one round-trip and one parse. So the
    saving is a single parse, on the Tier 0 path — and on the GROBID path the body is new
    work the pipeline did not previously do at all, because nothing in production built an
    artifact before this phase.
    """
    return _extract(path, engine, with_body=True)


def extract_references(path: Path, engine: str = ENGINE_AUTO) -> list[RawReferenceEntry]:
    """End-to-end: read a PDF/DOCX and split it into individual reference entries.

    For a PDF, tries GROBID (Tier 1: docs/project-plan.md measured it with
    substantially better DOI recall/precision than the Tier 0 regex splitter) first,
    falling back to the Tier 0 text-splitting path below whenever GROBID isn't
    reachable, fails on this particular document, or comes back with zero references
    — a 200 response with no parseable <biblStruct> entries (e.g. a scanned/OCR-hostile
    PDF, or a layout GROBID's segmentation model misses) isn't a GrobidUnavailableError,
    but treating it as "processed, zero references" instead of falling back would
    silently defeat the whole point of the fallback: GROBID is a best-effort upgrade,
    not a hard dependency, so a document still gets a real chance at extraction either
    way. DOCX always uses Tier 0 directly since GROBID has no DOCX support, regardless
    of engine.

    engine (see engine_status.ENGINE_*) lets a caller pin the extraction path instead
    of the default auto-detect: ENGINE_ANCHOR skips GROBID outright (e.g. to compare
    Tier 0 output against GROBID's, or when GROBID is known-unreliable for a given
    document); ENGINE_GROBID still falls back to Tier 0 on failure/zero-results —
    forcing a hard dependency on a best-effort upgrade would make an otherwise
    extractable document unreadable whenever the pinned engine has a bad day.

    The bibliography-only view of extract_document, and deliberately still the cheaper
    one: a caller that only wants references reads no body text and, on the GROBID path,
    never touches pymupdf4llm at all. Every existing caller is on this path, and the
    plan's first Success Criterion is that their output does not move.
    """
    return _extract(path, engine, with_body=False).references


def _extract(path: Path, engine: str, *, with_body: bool) -> ExtractedDocument:
    """extract_references and extract_document, as one implementation.

    Written once rather than twice because the two share every decision that can go
    wrong — which tier to try, when to fall back, where the bibliography stops — and two
    copies of that is the divergence signals.py's docstring is about. `with_body` buys
    strictly more work and changes nothing about the references produced, which is what
    lets the byte-identical regression test compare one against the other directly.
    """
    references, citations = _grobid_pass(path, engine)
    if references and not with_body:
        return ExtractedDocument(references=references, citations=citations)

    try:
        raw_text = _read_source_text(path)
        full_text = normalize_markdown_text(raw_text) if path.suffix.lower() == ".pdf" else raw_text
    except Exception:
        # Reading the document is required for the references and optional for the body.
        # A PDF that GROBID parsed but pymupdf4llm cannot open — encrypted, truncated —
        # used to produce a complete reference check, because the GROBID path never
        # touched pymupdf4llm at all. Letting that failure through now would take the
        # check with it, which is exactly what the plan's first Success Criterion forbids:
        # the reference output does not move when a citation input is unavailable.
        if not references:
            raise
        return ExtractedDocument(references=references, citations=citations)

    bibliography_text: str | None = None
    try:
        bibliography_text = find_bibliography_section(full_text)
    except NoBibliographySectionError:
        # Only fatal when this document's references depend on it. GROBID having already
        # produced them means the body is still worth reading — a manuscript whose
        # reference list this project cannot locate is exactly the one whose in-text
        # citations are worth showing.
        if not references:
            raise

    if not references and bibliography_text is not None:
        references = split_bibliography_block(bibliography_text)
    if not with_body:
        return ExtractedDocument(references=references, citations=citations)

    audit = _audit(references, bibliography_text)

    try:
        artifact = _artifact_from(raw_text, path.suffix.lower())
        profile = compute_style_profile(bibliography_text) if bibliography_text is not None else None
    except Exception:
        # Same rule one step later: by here the references exist and nothing the body
        # produces may endanger them. A caller that gets no artifact is told so by the
        # field being None, which `_match_citations` already reports as "the citation
        # search did not run" rather than as "this manuscript cites nothing".
        return ExtractedDocument(references=references, citations=citations, audit=audit)

    return ExtractedDocument(
        references=references, citations=citations, artifact=artifact, profile=profile, audit=audit
    )


def _audit(references: list[RawReferenceEntry], bibliography_text: str | None) -> ReferenceListAudit:
    """Compare the extracted list against the printed one.

    Two readers, one comparison. It is a report for the reviewer — how the extracted list
    differs from the page — and it is what a numbered marker resolves through, since until
    it existed "[12]" could only be answered by an entry's position in a list nothing had
    established was whole (see `citation_matching._reference_index`).

    Never allowed to fail the extraction. The reference list is the check's product and
    this is commentary on it: a bibliography whose shape defeats the alignment leaves the
    references exactly as they were and the audit unread, which is what a reader of
    `ReferenceListAudit.read` is told.
    """
    if bibliography_text is None:
        return NOT_AUDITED
    try:
        return audit_reference_list(references, bibliography_text)
    except Exception:
        return NOT_AUDITED


def _grobid_pass(path: Path, engine: str) -> tuple[list[RawReferenceEntry], tuple[CitationContext, ...]]:
    """Tier 1's half of extraction: the references and the in-text citations of them.

    Zero references is a fallback case rather than a result, exactly as before — a 200
    response with nothing parseable in it is not a GrobidUnavailableError, and treating it
    as "processed, zero references" would defeat the fallback. The citations go with them:
    a document whose reference list came from Tier 0 has no TEI ids to resolve against, so
    carrying GROBID's markers forward would hand the matcher targets pointing into a list
    that is not the one it holds.
    """
    use_grobid = path.suffix.lower() == ".pdf" and engine != ENGINE_ANCHOR and (
        is_grobid_available() if engine != ENGINE_GROBID else True
    )
    if not use_grobid:
        return [], ()
    try:
        extracted = extract_document_via_grobid(path)
    except GrobidUnavailableError:
        return [], ()
    if not extracted.references:
        return [], ()
    return extracted.references, tuple(extracted.citations)


def _read_source_text(path: Path) -> str:
    """The document as its parser renders it, before any normalization.

    Markdown for a PDF and joined paragraph text for a DOCX — the same two strings
    build_document_artifact keeps as the far end of its source-anchor mapping, read here
    so one extraction serves both the bibliography split and the artifact.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_markdown(path)
    if suffix == ".docx":
        return _extract_docx_text(path)
    raise ValueError(f"Unsupported file type: {suffix!r} (expected .pdf or .docx)")


def _docx_artifact_lines(raw_text: str) -> list[str]:
    """DOCX lines with their paragraph boundaries made explicit.

    _extract_docx_text joins python-docx's paragraphs with a single newline, so an
    ordinary DOCX — consecutive non-empty paragraphs, no blank ones — arrives as a run
    of adjacent non-blank lines with nothing marking where one paragraph ends.
    DocumentArtifact.from_lines separates paragraphs on blank lines, the convention
    pymupdf4llm's Markdown output already follows, so without this every body paragraph
    in a DOCX collapses into a single Paragraph: a citation in the tenth paragraph
    reports the first paragraph's source_line, and expanding its context hands the
    reviewer the whole body instead of the passage they asked for — the artifact's
    paragraph is what bounds that disclosure.

    A python-docx paragraph is a real paragraph, so the separator goes between every
    pair rather than being inferred from anything. Separators are tagged with no origin
    because they correspond to nothing in `raw_text`: that stays the unmodified join, so
    every other line's index keeps pointing at the line it actually came from.
    """
    lines: list[str] = []
    for i, line in enumerate(raw_text.splitlines()):
        if lines:
            lines.append(SourceLine("", None))
        lines.append(SourceLine(line, i))
    return lines


def build_document_artifact(path: Path) -> DocumentArtifact:
    """Read a PDF/DOCX and return its body text with source anchors.

    The companion to extract_references, not a replacement for it: that function's
    bibliography-only contract is unchanged and its callers are untouched, because a
    check must go on producing exactly today's reference output whether or not
    anything downstream ever asks for an artifact (see the citation plan's Success
    Criteria). The two share the same extraction and the same bibliography boundary,
    so they cannot disagree about where the reference list is — but nothing here can
    make extract_references behave differently.

    A document with no recognizable bibliography heading yields an artifact over the
    whole text rather than an error. extract_references is the caller entitled to fail
    on that; body text is still body text, and a manuscript whose reference list this
    project can't find is exactly the one whose in-text citations are worth showing.
    """
    return _artifact_from(_read_source_text(path), path.suffix.lower())


def _artifact_from(raw_text: str, suffix: str) -> DocumentArtifact:
    """build_document_artifact, given text something else already extracted.

    Split out so `extract_document` can build an artifact from the same extraction it
    split the bibliography out of, instead of parsing the PDF a second time — the cost
    Phase 1 recorded as something Phase 4 would otherwise pay on every check.
    """
    if suffix == ".pdf":
        lines: list[str] = list(_normalize_markdown_lines(raw_text.splitlines()))
        parser = "pymupdf4llm"
    elif suffix == ".docx":
        lines = _docx_artifact_lines(raw_text)
        parser = "python-docx"
    else:
        raise ValueError(f"Unsupported file type: {suffix!r} (expected .pdf or .docx)")

    bounds = _bibliography_bounds(lines)
    if bounds is None:
        body_lines = lines
    else:
        heading_idx, end_idx = bounds
        body_lines = lines[:heading_idx] + lines[end_idx:]
    return DocumentArtifact.from_lines(body_lines, raw_text=raw_text, parser=parser)
